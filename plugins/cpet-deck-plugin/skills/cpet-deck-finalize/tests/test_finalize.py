"""Self-contained tests for anim.py, derive.py and deck_guard.py (synthetic deck, no project data).

    python test_finalize.py
"""
import json
import shutil
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))
import anim  # noqa: E402
import deck_guard as g  # noqa: E402
import derive  # noqa: E402
from pptx_xml import Package  # noqa: E402

results = []


def check(name, cond, detail=""):
    results.append(cond)
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  ({detail})" if detail and not cond else ""))


def make_deck(path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, (name, text) in enumerate([("EntryPrompt", "Why?"), ("EntryPrompt2", "Because."),
                                      ("Card A", "first card"), ("Card B", "second card"),
                                      ("Takeaway", "conclusion")]):
        tb = s.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(6), Inches(0.6))
        tb.name = name
        tb.text_frame.text = text
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    s.notes_slide.notes_text_frame.text = "old notes"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "static"
    prs.save(path)


SPEC = {"slides": {"1": [
    [{"shape": "EntryPrompt2", "effect": "appear"}],
    [{"shape": "EntryPrompt", "effect": "exit"}, {"shape": "EntryPrompt2", "effect": "exit"},
     {"shape": "Card A", "effect": "appear"}],
    [{"shape": "Card B", "effect": "appear"}],
    [{"shape": "Takeaway", "effect": "appear"}],
]}}


def main():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        src = d / "src.pptx"
        make_deck(src)
        spec = d / "spec.json"
        spec.write_text(json.dumps(SPEC), encoding="utf-8")

        # --- anim inject / check / extract round trip
        final = d / "Final.pptx"
        check("inject returns 0 (lint clean)", anim.main(["inject", str(src), str(spec), str(final)]) == 0)
        pkg = Package(final)
        clicks = anim.read_clicks(pkg.tree(pkg.slides()[0]))
        check("4 clicks written", len(clicks) == 4, len(clicks))
        grp = {(e["spid"], e["grpId"]) for c in clicks for e in c}
        check("EntryPrompt2 appear/exit use distinct grpIds", len(grp) == sum(len(c) for c in clicks))
        spec2 = d / "spec2.json"
        anim.main(["extract", str(final), str(spec2)])
        check("extract reproduces the spec", json.loads(spec2.read_text(encoding="utf-8")) == SPEC)

        # --- D53 regression: shared grpId must be flagged
        bad = Package(final)
        t = bad.tree(bad.slides()[0])
        for c in t.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cTn"):
            if c.get("presetClass"):
                c.set("grpId", "0")
        bad.save(d / "bad.pptx")
        check("shared grpId flagged by check", anim.main(["check", str(d / "bad.pptx")]) == 1)

        # --- derive noanim + verify
        noanim = d / "NoAnim.pptx"
        derive.main(["noanim", str(final), str(noanim)])
        p = Presentation(noanim)
        names = [s.name for s in p.slides[0].shapes]
        check("noanim drops EntryPrompt* shapes", not any(n.startswith("EntryPrompt") for n in names), names)
        check("noanim has no timing", not anim.read_clicks(Package(noanim).tree(Package(noanim).slides()[0])))
        check("verify-noanim matches", derive.main(["verify-noanim", str(final), str(noanim)]) == 0)
        p.slides[0].shapes[0].text_frame.text = "drift"
        p.save(d / "NoAnim_drift.pptx")
        check("verify-noanim detects drift", derive.main(["verify-noanim", str(final), str(d / "NoAnim_drift.pptx")]) == 1)

        # --- notes markup
        notes = d / "notes.json"
        notes.write_text(json.dumps({"1": "[0] Say **this** and __the__ end.\n[1] Next."}), encoding="utf-8")
        derive.main(["notes", str(final), str(notes), str(d / "Final_n.pptx")])
        tf = Presentation(d / "Final_n.pptx").slides[0].notes_slide.notes_text_frame
        runs = [(r.text, r.font.bold, r.font.underline) for r in tf.paragraphs[0].runs]
        check("notes bold + thee underline", ("this", True, None) in runs and ("the", None, True) in runs, runs)
        check("notes paragraphs", len(tf.paragraphs) == 2)

        # --- guard
        man = str(d / "deck_manifest.json")
        g.main(["register", str(final), "--manifest", man])
        g.main(["register", str(noanim), "--manifest", man, "--role", "derived", "--from", str(final)])
        check("guard clean", g.main(["check", str(final), "--manifest", man]) == 0)
        p = Presentation(final)
        sh = p.slides[0].shapes[2]
        sh.top = Emu(sh.top + 18288)
        p.save(final)
        check("guard detects moved shape", g.main(["check", str(final), "--manifest", man]) == 1)
        g.main(["register", str(final), "--manifest", man])
        check("stale detects derived deck", g.main(["stale", "--manifest", man]) == 1)
        arch = d / "99_archive"
        g.main(["snapshot", str(final), "--archive", str(arch), "--action", "p01_edit"])
        g.main(["snapshot", str(final), "--archive", str(arch), "--action", "p01_edit"])
        folders = sorted(x.name for x in arch.iterdir())
        check("snapshots get unique folders with README",
              len(folders) == 2 and all((arch / f / "README.md").exists() for f in folders), folders)
        owner = final.with_name("~$" + final.name)
        owner.write_bytes(b"x")
        check("guard reports lock", g.main(["check", str(final), "--manifest", man]) == 3)
    failed = results.count(False)
    print("all passed" if not failed else f"{failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
