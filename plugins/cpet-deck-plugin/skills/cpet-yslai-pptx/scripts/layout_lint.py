#!/usr/bin/env python3
"""layout_lint.py — 對產出的 .pptx 做版面幾何與樣式檢查（v2.9 第二批）。

**本腳本補的是像素 QA 的根本缺口。** `qa_check.py` 只做像素色塊分布統計，
**不做文字框幾何碰撞偵測**——v2.8 開發過程中三個真實缺陷全部是它判 PASS、
由人眼逐張目視才抓到的：

1. `table` 頁的 note 與表格最後一列重疊（提案 A；同一份簡報三張表都踩到）
2. figure band 內有／無 `source` 的圖高度參差，下緣不齊
3. 封面資訊區被推出投影片左緣外、文字左半被裁掉

三者的共通點是**幾何錯位**：形狀的座標與尺寸本身就錯了，但顏色分布仍正常，
所以像素統計看不出來。本腳本直接讀 pptx 的形狀座標判斷，不經渲染。

判定原則（重疊）：**完全包含＝合法**（卡片包住其內文字框、圖形包住標籤都是正常的
容器關係），**部分重疊＝可疑**。圖片與表格是「實心」的，任何部分重疊都會遮住內容，
故門檻設低；兩個文字框則常帶留白，需達一定比例才判定。

用法：
    python layout_lint.py deck.pptx [--strict] [--quiet]
"""
import argparse
import re
import sys
import zipfile
from collections import defaultdict

from pptx import Presentation

# 主控台編碼：Windows 繁中環境的管線 stdout 預設 cp950，印「≤」「✔」等字元會 UnicodeEncodeError 崩潰。
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")
    except Exception:
        pass

EMU_IN = 914400
SLIDE_W, SLIDE_H = 13.333, 7.5

HEADER_BAR_BOTTOM = 0.28   # 頂部 header bar 下緣
MASTER_RULE_Y = 7.25       # 底部藍線
BODY_TOP = 0.96            # 標題 placeholder 下緣
BODY_BOTTOM = 7.02         # 內容安全下限
TOL = 0.02                 # 幾何容差（吋），避免浮點與四捨五入造成假警報

# 結構色（不計入語意色數）。取自 build_deck 的 C_* 常數與模板母片。
STRUCTURAL = {
    "0000FF", "1F00FF",
    "14274E", "003781",
    "1F2937", "1F1F6E", "000000", "FFFFFF",
    "6B7280", "595959", "A6A6A6",
    "F5F7FA", "D6DEE9", "DCE6F1", "8EA9C9",
    # v2.11：重點字 token（行內重點紅、結論橫條黃字）屬「強調」不屬語意分類，不計入色數
    "C00000", "FFD54F",
}
# v2.11：淡底卡／泳道的底色、橙的暗階，與本色同屬一個語意色，計數時歸併
TINT_FAMILY = {"E3F0F9": "0072B2", "FBEFD8": "E69F00", "8A5A00": "E69F00",
               "FBE9DD": "D55E00", "E0F4EE": "009E73"}
MIN_AUX_PT = 10.5          # 輔助文字的絕對下限


def apply_profile(path):
    """官方模板 profile：版面保留區與模板自帶顏色依 profile 調整（預設值對應實驗室模板）。"""
    global HEADER_BAR_BOTTOM, MASTER_RULE_Y, BODY_TOP, BODY_BOTTOM
    import json
    with open(path, encoding="utf-8") as f:
        prof = json.load(f)
    lint_spec = prof.get("lint", {})
    body = prof.get("body", {})
    HEADER_BAR_BOTTOM = float(lint_spec.get("header_bottom", 0.0))
    MASTER_RULE_Y = float(lint_spec.get("rule_y", SLIDE_H))
    BODY_TOP = float(body.get("top", BODY_TOP))
    BODY_BOTTOM = float(body.get("bottom", BODY_BOTTOM))
    STRUCTURAL.update(c.upper().lstrip("#") for c in lint_spec.get("structural_colors", []))

PIC, TABLE, TEXTBOX, AUTOSHAPE, PLACEHOLDER = 13, 19, 17, 1, 14


class Report:
    def __init__(self):
        self.errors, self.warnings = [], []

    def err(self, w, m):
        self.errors.append((w, m))

    def warn(self, w, m):
        self.warnings.append((w, m))


def _box(sh):
    try:
        return (sh.left / EMU_IN, sh.top / EMU_IN,
                (sh.left + sh.width) / EMU_IN, (sh.top + sh.height) / EMU_IN)
    except (TypeError, AttributeError):
        return None


def _inter(a, b):
    return (max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
            * max(0.0, min(a[3], b[3]) - max(a[1], b[1])))


def _contains(outer, inner, tol=TOL):
    return (outer[0] - tol <= inner[0] and outer[1] - tol <= inner[1]
            and outer[2] + tol >= inner[2] and outer[3] + tol >= inner[3])


def _has_text(sh):
    return sh.has_text_frame and sh.text_frame.text.strip() != ""


def _stype(sh):
    try:
        return int(sh.shape_type)
    except (TypeError, ValueError):
        return 0


def _label(sh):
    t = sh.text_frame.text.strip().replace("\n", " ")[:18] if _has_text(sh) else ""
    kind = {PIC: "圖片", TABLE: "表格", TEXTBOX: "文字框",
            AUTOSHAPE: "圖形", PLACEHOLDER: "版面配置區"}.get(_stype(sh), "形狀")
    return f"{kind}{('「' + t + '」') if t else ''}"


def _occupancy(boxes, x0, y0, x1, y1, nx=180, ny=100):
    """網格取樣估算聯集佔用比——比解析式聯集簡單，對本用途精度足夠。"""
    if not boxes or x1 <= x0 or y1 <= y0:
        return 0.0
    hit = 0
    dx, dy = (x1 - x0) / nx, (y1 - y0) / ny
    for i in range(nx):
        cx = x0 + (i + 0.5) * dx
        for j in range(ny):
            cy = y0 + (j + 0.5) * dy
            for b in boxes:
                if b[0] <= cx <= b[2] and b[1] <= cy <= b[3]:
                    hit += 1
                    break
    return hit / (nx * ny)


def lint(pptx_path):
    rep = Report()
    prs = Presentation(pptx_path)
    z = zipfile.ZipFile(pptx_path)
    names = sorted([n for n in z.namelist()
                    if n.startswith("ppt/slides/slide") and n.endswith(".xml")],
                   key=lambda n: int("".join(filter(str.isdigit, n.split("/")[-1]))))

    for idx, slide in enumerate(prs.slides, 1):
        where = f"第 {idx} 張"
        shapes = [s for s in slide.shapes if _box(s)]
        boxes = {id(s): _box(s) for s in shapes}

        # ── 1. 超出邊界／侵入母片保留區 ──
        for s in shapes:
            b = boxes[id(s)]
            if b[0] < -TOL or b[1] < -TOL or b[2] > SLIDE_W + TOL or b[3] > SLIDE_H + TOL:
                rep.err(where, f"{_label(s)} 超出投影片邊界："
                               f"({b[0]:.2f},{b[1]:.2f})–({b[2]:.2f},{b[3]:.2f})")
                continue
            if _has_text(s):
                if b[3] > MASTER_RULE_Y + TOL:
                    rep.err(where, f"{_label(s)} 下緣 {b[3]:.2f}in 越過母片底線 "
                                   f"{MASTER_RULE_Y}in——文字會被線壓住")
                if b[1] < HEADER_BAR_BOTTOM - TOL:
                    rep.err(where, f"{_label(s)} 上緣 {b[1]:.2f}in 侵入頂部 header bar")

        # ── 2. 實質重疊 ──
        solid = [s for s in shapes if _stype(s) in (PIC, TABLE) or _has_text(s)]
        for i in range(len(solid)):
            for j in range(i + 1, len(solid)):
                a, b = solid[i], solid[j]
                ba, bb = boxes[id(a)], boxes[id(b)]
                if _inter(ba, bb) <= 0:
                    continue
                # 容器豁免只適用於「真的是容器」的形狀：圓角矩形卡片、takeaway 橫幅、
                # 徽章、版面配置區。**表格與圖片不是容器**——文字框落在表格的 bbox 內
                # 一定是壓在資料上，不能豁免（回測時第一版誤放過了 table note 重疊）。
                a_container = _stype(a) not in (PIC, TABLE)
                b_container = _stype(b) not in (PIC, TABLE)
                if (_contains(ba, bb) and a_container) or (_contains(bb, ba) and b_container):
                    continue
                area = min((ba[2] - ba[0]) * (ba[3] - ba[1]),
                           (bb[2] - bb[0]) * (bb[3] - bb[1]))
                if area <= 0:
                    continue
                ratio = _inter(ba, bb) / area
                hard = _stype(a) in (PIC, TABLE) or _stype(b) in (PIC, TABLE)
                if (hard and ratio > 0.02) or ratio > 0.15:
                    lvl = rep.err if (hard or ratio > 0.35) else rep.warn
                    lvl(where, f"{_label(a)} 與 {_label(b)} 部分重疊"
                               f"（重疊佔較小者 {ratio:.0%}）")

        # ── 2b. 表格內容高度 vs 宣告高度 ──
        # 為何需要這條：PowerPoint 的表格列高會依內容自動撐開，**宣告高度 ≠ 實際渲染
        # 高度**。提案 A 的原始 bug 正是如此——note 畫在「宣告的」表格下緣之下，宣告
        # 幾何看起來沒有重疊，渲染後卻壓在最後一列上。因此單看形狀座標抓不到，必須
        # 另外估算表格內容所需的高度。v2.8 起建構器會把算好的列高明確寫回表格，
        # 這條檢查即是驗證那個機制是否失效。
        for s in shapes:
            if _stype(s) != TABLE or not s.has_table:
                continue
            tbl = s.table
            bx = boxes[id(s)]
            declared = bx[3] - bx[1]
            col_w = (bx[2] - bx[0]) / max(1, len(tbl.columns)) - 0.18
            need = 0.0
            for row in tbl.rows:
                pt = 20.0
                lines = 1
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        for r in para.runs:
                            if r.font.size:
                                pt = r.font.size.pt
                    w = sum(1.0 if ord(c) > 0x2E7F else 0.5 for c in cell.text)
                    if col_w > 0:
                        lines = max(lines, max(1, int(w * pt / 72.0 / col_w + 0.999)))
                need += max(pt * 1.75 / 72.0, lines * pt * 1.34 / 72.0 + 0.12)
            if need > declared + 0.10:
                rep.err(where, f"表格內容估算需 {need:.2f}in、宣告高度僅 {declared:.2f}in"
                               f"——PowerPoint 會自動撐開列高並壓到下方元素"
                               f"（這正是提案 A 的原始症狀，宣告座標看不出來）")

        # ── 3. 內容密度（留白）──
        content = [boxes[id(s)] for s in shapes
                   if _has_text(s) or _stype(s) in (PIC, TABLE)]
        content = [c for c in content if c[3] > BODY_TOP and c[1] < BODY_BOTTOM]
        if content:
            occ = _occupancy(content, 0.0, BODY_TOP, SLIDE_W, BODY_BOTTOM)
            has_fig = any(_stype(s) == PIC for s in shapes)
            cap = 0.85 if has_fig else 0.75
            if occ > cap:
                rep.warn(where, f"內容佔用內容區 {occ:.0%}"
                                f"（{'含圖' if has_fig else '純文字'}頁建議 ≤{cap:.0%}）"
                                f"——留白不足，建議拆頁或精簡")

        # ── 4. 同列圖片高度一致性 ──
        rows = defaultdict(list)
        for s in shapes:
            if _stype(s) == PIC:
                rows[round(boxes[id(s)][1], 1)].append(boxes[id(s)])
        for group in rows.values():
            if len(group) < 2:
                continue
            hs = [g[3] - g[1] for g in group]
            if max(hs) - min(hs) > 0.15:
                rep.warn(where, f"同列 {len(group)} 張圖高度不一致"
                                f"（{min(hs):.2f}–{max(hs):.2f}in）——下緣會參差")

        # ── 5. 字級下限 ──
        for s in shapes:
            if not _has_text(s):
                continue
            for para in s.text_frame.paragraphs:
                for r in para.runs:
                    if r.font.size is not None and r.font.size.pt < MIN_AUX_PT:
                        rep.warn(where, f"出現 {r.font.size.pt:g}pt 文字"
                                        f"（低於輔助文字下限 {MIN_AUX_PT:g}pt）："
                                        f"{r.text[:20]!r}")

        # ── 6. 語意色數 ──
        xml = z.read(names[idx - 1]).decode("utf-8", "ignore")
        used = {c.upper() for c in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml)}
        semantic = sorted({TINT_FAMILY.get(c, c) for c in used - STRUCTURAL})
        if len(semantic) > 3:
            rep.warn(where, f"語意色 {len(semantic)} 種（建議 ≤3，見 conventions §2.5.2）："
                            f"{'、'.join('#' + c for c in semantic[:6])}")
    return rep, len(prs.slides)


def print_report(rep, n, quiet=False):
    for w, m in rep.errors:
        print(f"ERROR  {w}\n       {m}")
    for w, m in rep.warnings:
        print(f"WARN   {w}\n       {m}")
    if rep.errors or rep.warnings:
        print()
    if rep.errors:
        print(f"✘ 版面幾何檢查：{len(rep.errors)} 個錯誤、{len(rep.warnings)} 則警告 — "
              f"錯誤代表形狀座標本身就不對（重疊／溢出），像素 QA 抓不到這類問題。")
    elif not quiet:
        print(f"✔ 版面幾何檢查通過（{n} 張，{len(rep.warnings)} 則警告）。"
              f"仍須逐張目視確認觀感。")


def main():
    ap = argparse.ArgumentParser(description="pptx 版面幾何與樣式檢查")
    ap.add_argument("pptx")
    ap.add_argument("--strict", action="store_true", help="警告也視為失敗")
    ap.add_argument("--quiet", action="store_true")
    ap.add_argument("--profile", help="官方模板 profile（與 build_deck.py --profile 相同的檔案）")
    a = ap.parse_args()
    if a.profile:
        apply_profile(a.profile)
    rep, n = lint(a.pptx)
    print_report(rep, n, quiet=a.quiet)
    if rep.errors or (a.strict and rep.warnings):
        sys.exit(1)


if __name__ == "__main__":
    main()
