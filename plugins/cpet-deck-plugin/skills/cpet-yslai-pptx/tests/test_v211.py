"""v2.11 tests: plan_lint rules and build_deck features (synthetic plans, lab template).

    python test_v211.py
Builds into a temp folder; needs python-pptx, lxml, Pillow and matplotlib (for equations).
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
import plan_lint  # noqa: E402

ok = []


def check(name, cond, detail=""):
    ok.append(bool(cond))
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  ({detail})" if not cond and detail else ""))


def lint(plan, d):
    p = Path(d) / "p.json"
    p.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    rep = plan_lint.lint(str(p))
    return [m for _, m in rep.errors], [m for _, m in rep.warnings]


def bands(*bs, **kw):
    return {"type": "bands", "title": "T", "bands": list(bs), **kw}


def main():
    with tempfile.TemporaryDirectory() as d:
        e, w = lint({"slides": [bands({"kind": "text", "bullets": ["a"]}, takeaway="x !!y")]}, d)
        check("lint: unpaired !! in takeaway is an error", any("沒有成對" in m for m in e), e)
        e, w = lint({"slides": [bands({"kind": "text", "bullets": ["a"]},
                                      takeaway="!!a!! !!b!! !!c!! !!d!!")]}, d)
        check("lint: >3 takeaway highlights warns", any("重點片語" in m for m in w), w)
        e, w = lint({"slides": [bands({"kind": "cards", "cards": [{"head": "h", "tint": "pink"}]})]}, d)
        check("lint: unknown tint is an error", any("tint" in m for m in e), e)
        e, w = lint({"slides": [bands({"kind": "cards", "cards": [
            {"head": "a", "tint": "blue"}, {"head": "b", "tint": "vermilion"}, {"head": "c", "tint": "green"}]})]}, d)
        check("lint: 4 semantic colours (with orange) warns", any("語意色" in m for m in w), w)
        e, w = lint({"meta": {"deck_type": "conference-oral"}, "slides": [bands(
            {"kind": "cards", "cards": [{"head": "a", "bullets": ["1", "2", "3", "4"]}]})]}, d)
        check("lint: conference-oral card >3 bullets warns", any("≤3 點" in m for m in w), w)
        e, w = lint({"slides": [bands({"kind": "chevrons", "accent": 9, "steps": [{"head": "a"}, {"head": "b"}]})]}, d)
        check("lint: chevrons accent out of range", any("accent" in m for m in e), e)
        e, w = lint({"slides": [bands({"kind": "lanes", "lanes": [{"label": "A", "items": [{"x": 1}]}]})]}, d)
        check("lint: lane item needs image or text", any("image 或 text" in m for m in e), e)
        e, w = lint({"meta": {"visual": "v2.11", "title_style": "two-level"},
                     "slides": [bands({"kind": "text", "bullets": ["a"]}, notes="n", backup=True)]}, d)
        check("lint: v2.11 keys pass lint cleanly", not e and not w, (e, w))

        # build
        plan = {"meta": {"language": "en", "title_style": "two-level"}, "slides": [
            {"type": "cover", "title_en": "Demo", "presenter": "X"},
            bands({"kind": "chevrons", "accent": 0, "steps": [{"head": "A", "body": "a"}, {"head": "B"}, {"head": "C"}]},
                  {"kind": "cards", "cols": 4, "cards": [{"head": "Tinted", "tint": "blue", "body": "with !!stress!!"},
                                                         {"kpi": "≈ 57.5 hours", "body": "narrow kpi"},
                                                         {"head": "c"}, {"head": "d"}]},
                  section="Method", takeaway="one !!two!! three", notes="[0] hello"),
            {"type": "equation", "section": "Method", "title": "Eq", "blocks": [{"latex": "a=b", "note": "n"}]},
            bands({"kind": "text", "bullets": ["backup page"]}, backup=True)]}
        pj = Path(d) / "plan.json"
        pj.write_text(json.dumps(plan), encoding="utf-8")
        out = Path(d) / "out.pptx"
        r = subprocess.run([sys.executable, str(SCRIPTS / "build_deck.py"), str(pj), str(out)],
                           capture_output=True, text=True, encoding="utf-8")
        check("build succeeds", r.returncode == 0 and out.exists(), r.stdout[-300:] + r.stderr[-300:])
        prs = Presentation(out)
        titles = [s.shapes.title.text_frame.text if s.shapes.title else "" for s in prs.slides]
        check("backup slide comes after closing", len(prs.slides) == 5 and "Backup" in titles[4]
              and "Thank" in " ".join(sh.text_frame.text for sh in prs.slides[3].shapes if sh.has_text_frame), titles)
        tr = [r for r in prs.slides[1].shapes.title.text_frame.paragraphs[0].runs]
        check("two-level title: section 32pt small caps + page 16pt bold",
              tr[0].font.size.pt == 32 and tr[0]._r.rPr.get("cap") == "small" and tr[1].font.size.pt == 16
              and tr[1].font.bold, [(x.text, x.font.size) for x in tr])
        runs = [r for sh in prs.slides[1].shapes if sh.has_text_frame
                for p in sh.text_frame.paragraphs for r in p.runs]
        check("takeaway highlight yellow +4pt", any(r.text == "two" and str(r.font.color.rgb) == "FFD54F"
                                                   and r.font.size.pt == 20 for r in runs))
        check("card stress red bold", any(r.text == "stress" and str(r.font.color.rgb) == "C00000"
                                          and r.font.bold for r in runs))
        texts = [sh.text_frame.text for sh in prs.slides[2].shapes if sh.has_text_frame]
        check("equation source not on slide (v2.11)", not any(t.startswith("LaTeX") for t in texts), texts)
        notes = prs.slides[2].notes_slide.notes_text_frame.text
        check("equation source in notes", "LaTeX" in notes and "a=b" in notes, notes)
        check("plan notes written", prs.slides[1].notes_slide.notes_text_frame.text == "[0] hello")
        kpi = [r for sh in prs.slides[1].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs
               for r in p.runs if r.text == "≈ 57.5 hours"]
        check("narrow KPI shrinks below 44pt", kpi and kpi[0].font.size.pt < 44, [k.font.size for k in kpi])

        # v2.8 visual keeps the old look
        plan["meta"]["visual"] = "v2.8"
        pj.write_text(json.dumps(plan), encoding="utf-8")
        out2 = Path(d) / "out28.pptx"
        subprocess.run([sys.executable, str(SCRIPTS / "build_deck.py"), str(pj), str(out2)], capture_output=True)
        texts = [sh.text_frame.text for sh in Presentation(out2).slides[2].shapes if sh.has_text_frame]
        check("visual v2.8: equation source stays on slide", any(t.startswith("LaTeX") for t in texts), texts)
    failed = ok.count(False)
    print("all passed" if not failed else f"{failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
