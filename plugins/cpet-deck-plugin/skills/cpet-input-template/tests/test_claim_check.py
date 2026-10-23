"""Self-contained tests for scripts/claim_check.py (synthetic numbers, no project data).

Reproduces the ECCE 2026 D49 pattern: a KPI tile shows the *command* value, a result card shows
the *measured* value for the same condition without saying so.   python test_claim_check.py
"""
import json
import sys
import tempfile
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import claim_check as cc  # noqa: E402

LEDGER = """# 數據查核表
| id | anchor | value | unit | kind | labels | source | pages | status |
|---|---|---|---|---|---|---|---|---|
| T9_cmd_new | 9 000 rpm | 40 | Nm | command | command;lut | synthetic | #1 #2 | 已查證 |
| T9_cmd_old | 9 000 rpm | 33 | Nm | command | command;lut | synthetic | #2 | 已查證 |
| T9_meas_new | 9 000 rpm | 42.5 | Nm | measured | measured | synthetic | #2 | 已查證 |
| T9_meas_old | 9 000 rpm | 36.1 | Nm | measured | measured | synthetic | #2 | 已查證 |
| X1 | 待填 | 待填 | s | 待填 |  | 待填 |  | 待查證 |
"""


def box(slide, x, y, w, h, text, pt=16):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.text = text
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(pt)


def deck(path, kpi_value, card_text):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)  # KPI tile: value box directly above its condition
    box(s, 1, 2, 2, 0.6, "200 Nm", 32); box(s, 1, 2.65, 2, 0.4, "up to 2500 rpm")
    box(s, 4, 2, 2, 0.6, kpi_value, 32); box(s, 4, 2.65, 2, 0.4, "at 9 000 rpm")
    s = prs.slides.add_slide(blank)  # result card
    box(s, 1, 1, 7, 1, card_text)
    s = prs.slides.add_slide(blank)  # spec list: a different quantity next to the anchor
    box(s, 1, 1, 7, 1, "Motor: 200 Nm / 9 000 rpm, 45 kW")
    prs.save(path)


CASES = [
    ("labelled", "40 Nm", "Torque at 9 000 rpm: command 40 vs 33 Nm, measured 42.5 vs 36.1 Nm", (0, 1)),
    ("unlabelled measured (D49)", "40 Nm", "Higher torque at 9 000 rpm: 42.5 vs 36.1 Nm", (0, 3)),
    ("wrong KPI value", "41 Nm", "Torque at 9 000 rpm: command 40 vs 33 Nm, measured 42.5 vs 36.1 Nm", (1, 0)),
]

results = []


def check(name, cond, detail=""):
    results.append(cond)
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  ({detail})" if not cond and detail else ""))


def main():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        ledger = d / "數據查核表.md"
        ledger.write_text(LEDGER, encoding="utf-8")
        rows, skipped = cc.parse_ledger(ledger)
        check("non-numeric row skipped, not fatal", len(rows) == 4 and skipped == ["X1"], (rows, skipped))
        for name, kpi, card, (exp_m, exp_u) in CASES:
            p = d / "t.pptx"
            deck(p, kpi, card)
            f = cc.check(rows, [p])
            m = sum(x[0] == "MISMATCH" for x in f)
            check(f"{name}: mismatch {exp_m}, unlabeled {exp_u}", (m, len(f) - m) == (exp_m, exp_u), f)

        # talk.json findings point at a page
        talk = {"schema": "cpet-talk/0.1", "pages": [{"n": 7, "lines": [[0, "At 9 000 rpm we get 41 Nm."]]}]}
        tj = d / "talk.json"; tj.write_text(json.dumps(talk), encoding="utf-8")
        f = cc.check(rows, [tj])
        check("talk.json mismatch located by page", f and f[0][0] == "MISMATCH" and "talk p7" in f[0][1], f)
        check("exit code 1 on mismatch", cc.main(["check", str(ledger), str(tj)]) == 1)

        # suggest
        tmpl = d / "input_template.md"
        tmpl.write_text("## 結果\n### Torque\nTorque 40 Nm at 9 000 rpm, error 4.20 % → 2.10 %\n"
                        "### Time\nOnly 120 s instead of 40 h; 3 steps.\n### Again\n40 Nm again\n",
                        encoding="utf-8")
        buf = StringIO()
        with redirect_stdout(buf):
            cc.main(["suggest", str(tmpl)])
        out = buf.getvalue()
        lines = [l for l in out.splitlines() if l.startswith("| C")]
        vals = [(l.split("|")[3].strip(), l.split("|")[4].strip()) for l in lines]
        check("suggest finds number+unit pairs",
              vals == [("40", "Nm"), ("9 000", "rpm"), ("4.20", "%"), ("2.10", "%"), ("120", "s"), ("40", "h")], vals)
        check("suggest merges repeated value across pages", "Torque；Again" in lines[0], lines[0])
        check("suggest ignores '3 steps'", ("3", "s") not in vals)
        sugg = d / "suggested.md"; sugg.write_text(out, encoding="utf-8")
        check("suggest output parses as a ledger (all rows 待填 anchor)", cc.parse_ledger(sugg)[0] != [])
    failed = results.count(False)
    print("all passed" if not failed else f"{failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
