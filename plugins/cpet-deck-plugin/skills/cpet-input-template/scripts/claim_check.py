"""claim_check.py — keep every number in a deck consistent with the 數據查核表 (claim ledger).

  suggest <input_template.md>                      list number+unit occurrences per page, as ledger
                                                   rows to start 數據查核表.md (kind / source left 待填)
  check   <數據查核表.md> <deck.pptx|talk.json|*.md|*.txt> [...] [--strict]
                                                   cross-check the deck, notes and script against it

Ledger = the first markdown table whose header contains `id | anchor | value | unit | kind`
(extra columns such as `labels`, `source`, `pages`, `status` are kept but not required):

    | id | anchor | value | unit | kind | labels | source | pages | status |
    |---|---|---|---|---|---|---|---|---|
    | T9_cmd | 9 000 rpm | 40 | Nm | command | command;LUT | LUT.mat Te_max(end) | #3 #7 | 已查證 |
    | T9_meas | 9 000 rpm | 42.5 | Nm | measured | measured | dyno log | #7 | 已查證 |

check, at every place an anchor appears (each text box, joined with its nearest neighbour box
above / to the left; notes; each line of a text file):
  * MISMATCH (error): a number with the claim's unit near the anchor is not a ledger value;
  * UNLABELED (warning): the ledger holds more than one kind for that anchor + unit (e.g. command
    and measured) and the number has none of its claim's label words just before it.
Rows whose value is not a number, or whose kind is blank / 待填, take part only partially
(non-numeric rows are skipped; unknown kinds do not trigger UNLABELED).

Limits: spoken-form numbers ("forty") are not parsed; anchors are plain text with
flexible spaces / thousand separators, not regexes.
Exit: 0 clean, 1 MISMATCH found (or any warning with --strict), 2 input error.
"""
from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from pathlib import Path

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass

THIN, NNBSP, NBHY, MINUS = chr(0x2009), chr(0x202F), chr(0x2011), chr(0x2212)
SEP = f"[ {THIN}{NNBSP},]"                      # thousand separators seen in slides
NUM = rf"[-{MINUS}]?\d{{1,3}}(?:{SEP}\d{{3}})+(?:\.\d+)?|[-{MINUS}]?\d+(?:\.\d+)?"
WINDOW = 80
LABEL_SPAN = 30
NEIGHBOUR_GAP_EMU = 274320  # 0.3 in: boxes further apart are separate statements
BEFORE_GAP = re.compile(r"\s*(?:at|@|,|in)?\s*", re.I)
UNKNOWN_KIND = {"", "待填", "?", "tbd"}
SUGGEST_UNITS = ["%-points", "%", "Nm", "rpm", "kW", "W", "kV", "V", "mA", "A", "kHz", "Hz",
                 "ms", "s", "min", "h", "PU", "mH", "Wb", "dB", "°C", "°"]


def norm_num(s: str) -> float:
    return float(re.sub(SEP, "", s).replace(MINUS, "-"))


def anchor_regex(anchor: str) -> re.Pattern:
    # merge thousand groups ("9 000" -> "9000"), then let digits take any separator
    merged = re.sub(rf"(?<=\d){SEP}(?=\d{{3}}\b)", "", anchor.strip())
    parts = []
    for tok in re.findall(r"\d+|[^\W\d]+|[^\w\s]", merged):
        parts.append(f"{SEP}?".join(tok) if tok.isdigit() else re.escape(tok))
    return re.compile(r"(?<![\d.])" + rf"[\s{THIN}{NNBSP}{NBHY}-]*".join(parts), re.I)


def value_regex(unit: str) -> re.Pattern:
    # "40 Nm", "40 vs 33 Nm", "40 / 33 Nm"
    return re.compile(rf"({NUM})(?:\s*(?:vs\.?|/|and)\s*({NUM}))?\s*{re.escape(unit)}(?![A-Za-z])")


# ------------------------------------------------------------------ ledger
def parse_ledger(path: Path) -> tuple[list[dict], list[str]]:
    rows, skipped, header = [], [], None
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip().startswith("|"):
            if header and (rows or skipped):
                break
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if header is None:
            low = [c.lower() for c in cells]
            if {"id", "anchor", "value", "unit", "kind"} <= set(low):
                header = low
            continue
        if set("".join(cells)) <= set("-: "):
            continue
        row = dict(zip(header, cells))
        try:
            row["value"] = norm_num(row["value"])
        except ValueError:
            skipped.append(row.get("id", "?"))
            continue
        row["labels"] = [w.strip().lower() for w in row.get("labels", "").split(";") if w.strip()]
        rows.append(row)
    if header is None:
        raise ValueError(f"no ledger table with id|anchor|value|unit|kind in {path}")
    return rows, skipped


# ------------------------------------------------------------------ sources
def pptx_texts(path: Path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes):
        for sh in shapes:
            box = (sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0)
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            elif sh.has_text_frame and sh.text_frame.text.strip():
                yield box, sh.text_frame.text.replace("\n", " ")
            elif getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    yield box, " | ".join(c.text for c in r.cells)

    def neighbour(items, k, side):
        """Nearest text box directly above (x-overlap) or to the left (y-overlap)."""
        (x, y, w, h), _ = items[k]
        best, dist = None, None
        for j, ((x2, y2, w2, h2), t2) in enumerate(items):
            if j == k:
                continue
            if side == "above" and y2 + h2 <= y + h * 0.5 and min(x + w, x2 + w2) - max(x, x2) > 0:
                d = y - (y2 + h2)
            elif side == "left" and x2 + w2 <= x + w * 0.5 and min(y + h, y2 + h2) - max(y, y2) > 0:
                d = x - (x2 + w2)
            else:
                continue
            if d <= NEIGHBOUR_GAP_EMU and (dist is None or d < dist):
                best, dist = t2, d
        return best

    # A value and its condition often sit in separate boxes (KPI tile: "40 Nm" above
    # "at 9 000 rpm"), so each box is checked alone and joined with its nearest neighbours.
    prs = Presentation(path)
    for i, s in enumerate(prs.slides, 1):
        items = list(walk(s.shapes))
        for k, (_, text) in enumerate(items):
            yield f"slide {i}", text
            for side in ("above", "left"):
                nb = neighbour(items, k, side)
                if nb:
                    yield f"slide {i}", nb + "  " + text
        if s.has_notes_slide:
            yield f"notes {i}", s.notes_slide.notes_text_frame.text.replace("\n", " ")


def file_texts(path: Path):
    if path.suffix.lower() == ".pptx":
        yield from pptx_texts(path)
        return
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() == ".json":
        data = json.loads(text)
        if data.get("schema") == "cpet-talk/0.1":  # per page, so findings point at a page
            for p in data.get("pages", []):
                yield f"talk p{p['n']}", " ".join(t for _, t in p["lines"])
            for b in data.get("backup", []):
                yield f"talk p{b['n']}", " ".join(q + " " + a for q, a in b["qa"])
            for o in data.get("other", []):
                yield "talk other", f"{o['q']} {o['a']} {o.get('note', '')}"
        else:
            yield path.name, json.dumps(data, ensure_ascii=False)
        return
    for n, line in enumerate(text.splitlines(), 1):
        yield f"line {n}", line


# ------------------------------------------------------------------ check
def check(ledger: list[dict], sources: list[Path]):
    groups = defaultdict(list)  # (anchor, unit) -> claims
    for c in ledger:
        groups[(c["anchor"], c["unit"])].append(c)
    findings = []
    for src in sources:
        for loc, text in file_texts(src):
            for (anchor, unit), claims in groups.items():
                kinds = {c["kind"] for c in claims if c["kind"].strip().lower() not in UNKNOWN_KIND}
                vre = value_regex(unit)
                for am in anchor_regex(anchor).finditer(text):
                    hits = list(vre.finditer(text, am.end(), min(len(text), am.end() + WINDOW)))
                    if not hits:
                        # "40 Nm at 9 000 rpm": accept a value just before the anchor,
                        # but not "200 Nm / 9 000 rpm" (a different quantity in a spec list)
                        before = list(vre.finditer(text, max(0, am.start() - WINDOW), am.start()))
                        if before and BEFORE_GAP.fullmatch(text[before[-1].end():am.start()]):
                            hits = before[-1:]
                    for m in hits:
                        for g in (1, 2):
                            if not m.group(g):
                                continue
                            v = norm_num(m.group(g))
                            match = [c for c in claims if abs(c["value"] - v) < 1e-9]
                            lo, hi = min(am.start(), m.start()), max(am.end(), m.end())
                            snippet = text[max(0, lo - 10): hi + 10].strip()
                            where = f"{src.name} {loc}"
                            if not match:
                                findings.append(("MISMATCH", where, anchor,
                                                 f"{v:g} {unit} not in ledger {sorted(c['value'] for c in claims)}",
                                                 snippet))
                                continue
                            if len(kinds) > 1:
                                pre = text[max(0, m.start(g) - LABEL_SPAN): m.start(g)].lower()
                                labels = {w for c in match for w in c["labels"]}
                                if labels and not any(w in pre for w in labels):
                                    findings.append(("UNLABELED", where, anchor,
                                                     f"{v:g} {unit} is {'/'.join(sorted({c['kind'] for c in match}))} "
                                                     f"but no label {sorted(labels)} just before it", snippet))
    seen, unique = set(), []
    for f in findings:  # neighbour joins can report the same value twice
        if f[:4] not in seen:
            seen.add(f[:4])
            unique.append(f)
    return unique


def cmd_check(argv) -> int:
    strict = "--strict" in argv
    argv = [a for a in argv if a != "--strict"]
    if len(argv) < 2:
        print(__doc__)
        return 2
    try:
        ledger, skipped = parse_ledger(Path(argv[0]))
        findings = check(ledger, [Path(p) for p in argv[1:]])
    except (OSError, ValueError, KeyError) as e:
        print("error:", e, file=sys.stderr)
        return 2
    for level, loc, anchor, msg, snip in findings:
        print(f"{level:9} {loc:28} [{anchor}] {msg}\n          … {snip}")
    if skipped:
        print(f"skipped rows with non-numeric value: {skipped}")
    errs = sum(f[0] == "MISMATCH" for f in findings)
    warns = len(findings) - errs
    print(f"{len(ledger)} claims, {errs} mismatch, {warns} unlabeled")
    return 1 if errs or (strict and warns) else 0


# ------------------------------------------------------------------ suggest
def cmd_suggest(argv) -> int:
    if len(argv) != 1:
        print(__doc__)
        return 2
    text = Path(argv[0]).read_text(encoding="utf-8")
    units = "|".join(re.escape(u) for u in SUGGEST_UNITS)
    rx = re.compile(rf"({NUM})\s*({units})(?![A-Za-z])")
    page, seen, rows = "(前言)", set(), []
    for line in text.splitlines():
        h = re.match(r"^#{2,3}\s+(.*)", line)
        if h:
            page = h.group(1).strip()
            continue
        for m in rx.finditer(line):
            key = (m.group(1), m.group(2))
            ctx = line[max(0, m.start() - 25): m.end() + 25].strip().replace("|", "／")
            if key in seen:
                for r in rows:
                    if (r[1], r[2]) == key and page not in r[3]:
                        r[3].append(page)
                continue
            seen.add(key)
            rows.append([f"C{len(rows) + 1:02d}", m.group(1), m.group(2), [page], ctx])
    print("| id | anchor | value | unit | kind | labels | source | pages | status |")
    print("|---|---|---|---|---|---|---|---|---|")
    for cid, val, unit, pages, ctx in rows:
        print(f"| {cid} | 待填 | {val} | {unit} | 待填 |  | 待填 | {'；'.join(pages)} | 待查證 |")
    print(f"\n<!-- {len(rows)} 個數值。上下文（填 anchor 用）：")
    for cid, val, unit, pages, ctx in rows:
        print(f"{cid}: …{ctx}…")
    print("-->")
    return 0


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv and argv[0] == "suggest":
        return cmd_suggest(argv[1:])
    if argv and argv[0] == "check":
        return cmd_check(argv[1:])
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main())
