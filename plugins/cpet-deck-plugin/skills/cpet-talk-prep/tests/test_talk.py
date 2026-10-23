"""Self-contained tests for talk.py (synthetic talk, no project data).

    python test_talk.py
Also exercises the hand-off to cpet-deck-finalize (derive.py notes) when that skill sits next to this one.
"""
import json
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))
import talk  # noqa: E402

TALK = {
    "schema": "cpet-talk/0.1",
    "meta": {"title": "Demo", "wpm": 90, "wpm_alt": 105, "click_sec": 0.7, "page_sec": 1.0,
             "core_pages": [2], "checkpoints": [2], "limit_sec": 60,
             "spoken": {"LUTs": "lookup tables", "LUT": "lookup table"}},
    "thee_before": ["input", "mtpa"],
    "pages": [
        {"n": 1, "title": "Title", "hint": "開場", "lines": [[0, "Good morning."], [0, "I build the **LUT** today."]]},
        {"n": 2, "title": "Method", "hint": "核心頁", "notes_pre": ["【發音】LUT 唸 lookup table"],
         "lines": [[0, "Look at the input."], [1, "Then the MTPA curve, and the input-output map."],
                   [2, "Two LUTs are built."], [0, "No extra click here."]]},
    ],
    "backup": [{"n": 3, "title": "B1 · Limits", "hint": "限制\n第二行", "qa": [["Q1 · Why?", "Because of the **input**."]]}],
    "phrases": [["聽不清楚", "Could you repeat that?"]],
    "phrases_note": "原則：先講結論。",
    "other": [{"q": "Q2 · How?", "a": "By a sweep.", "note": "備註文字"}],
}

ok = []


def check(name, cond, detail=""):
    ok.append(cond)
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  ({detail})" if not cond and detail else ""))


def main():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        tj = d / "talk.json"
        tj.write_text(json.dumps(TALK, ensure_ascii=False), encoding="utf-8")
        t = talk.load(tj)

        rows, total = talk.timing_table(t)
        check("clicks per page", [r["clicks"] for r in rows] == [0, 2], rows)
        check("words counted with markup", rows[0]["words"] == 7, rows[0])
        check("timing within limit -> rc 0", talk.main(["timing", str(tj)]) == 0)

        payload = talk.notes_payload(t)
        p2 = payload["2"].split("\n")
        check("notes_pre first", p2[0].startswith("【發音】"))
        check("click markers", p2[2].startswith("[1] ") and p2[4] == "No extra click here.", p2)
        check("thee marks the listed word only", "__the__ input." in p2[1] and "__the__ MTPA" in p2[2]
              and "__the__ input-output" not in p2[2], p2)
        check("bold removed in notes", "**" not in payload["1"])
        b = payload["3"].split("\n")
        check("backup notes layout", b[0] == "【中文提示】限制 ／ 第二行" and b[2] == "Q: Q1 · Why?"
              and b[3] == "A: Because of __the__ input.", b)

        talk.main(["txt", str(tj), str(d / "txt"), "--prefix", "Demo"])
        p2t = (d / "txt" / "Demo_p02.txt").read_text(encoding="utf-8")
        check("txt: one paragraph per click, click-0 continues", p2t.count("\n\n") == 2
              and p2t.endswith("Two lookup tables are built. No extra click here."), repr(p2t))
        check("txt: LUT spoken", "lookup table" in (d / "txt" / "Demo_p01.txt").read_text(encoding="utf-8"))
        allt = (d / "txt" / "Demo_all.txt").read_text(encoding="utf-8")
        check("txt: all has headings and core mark", "### Slide 2: Method  ★ 核心頁" in allt)

        new = json.loads(json.dumps(TALK)); new["pages"][1]["lines"][2][1] = "Three LUTs are built."
        nj = d / "new.json"; nj.write_text(json.dumps(new, ensure_ascii=False), encoding="utf-8")
        check("diff reports change (rc 1)", talk.main(["diff", str(tj), str(nj)]) == 1)
        check("diff no change (rc 0)", talk.main(["diff", str(tj), str(tj)]) == 0)

        out = d / "speech.pptx"
        talk.main(["deck", str(tj), str(out)])
        prs = Presentation(out)
        texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
        check("deck: pages + backup + phrases + other", len(prs.slides) == 5, len(prs.slides))
        check("deck: timing in header", "click" in texts[1] and "words" in texts[1])
        runs = [r for sh in prs.slides[1].shapes if sh.has_text_frame
                for p in sh.text_frame.paragraphs for r in p.runs]
        check("deck: thee run underlined blue",
              any(r.text == "the" and r.font.underline and str(r.font.color.rgb) == "0070C0" for r in runs))
        check("deck: phrases footnote", "原則：先講結論。" in texts[3])

        # hand-off to cpet-deck-finalize
        fin = HERE.parent.parent / "cpet-deck-finalize" / "scripts"
        if fin.exists():
            sys.path.insert(0, str(fin))
            import derive
            deck = Presentation(); deck.slide_width = Inches(13.333)
            for _ in range(3):
                s = deck.slides.add_slide(deck.slide_layouts[6]); s.notes_slide.notes_text_frame.text = "old"
            deck.save(d / "deck.pptx")
            talk.main(["notes", str(tj), str(d / "notes.json")])
            rc = derive.main(["notes", str(d / "deck.pptx"), str(d / "notes.json"), str(d / "deck_n.pptx")])
            tf = Presentation(d / "deck_n.pptx").slides[1].notes_slide.notes_text_frame
            check("finalize hand-off writes notes", rc == 0 and tf.paragraphs[2].text.startswith("[1] Then the MTPA"))
        else:
            print("SKIP  finalize hand-off (cpet-deck-finalize not next to this skill)")
    failed = ok.count(False)
    print("all passed" if not failed else f"{failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
