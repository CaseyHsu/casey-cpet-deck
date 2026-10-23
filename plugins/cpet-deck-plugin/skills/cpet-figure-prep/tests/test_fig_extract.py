"""Self-contained tests for scripts/fig_extract.py (synthetic .fig files, no project data).

    python test_fig_extract.py
Needs numpy, scipy, matplotlib, Pillow, python-pptx.
"""
import sys
import tempfile
from pathlib import Path

import numpy as np
import scipy.io as sio
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import fig_extract as fe  # noqa: E402

ok = []


def check(name, cond, detail=""):
    ok.append(bool(cond))
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  ({detail})" if not cond and detail else ""))


def node(typ, props, children=()):
    return {"type": typ, "handle": 1.0, "properties": props,
            "children": np.array(list(children), dtype=object) if children else np.zeros((0, 0))}


def classic_fig(path):
    """Mimic the MAT v5 layout of an old-style .fig: figure > axes > lines + label texts."""
    x = np.linspace(2000, 9000, 50)
    l1 = node("graph2d.lineseries", {"XData": x, "YData": 300 - x / 50, "DisplayName": "FW trajectory",
                                      "Color": np.array([0, 0.447, 0.741]), "LineStyle": "-", "Marker": "o"})
    l2 = node("graph2d.lineseries", {"XData": x, "YData": 250 - x / 80, "DisplayName": "MTPV trajectory",
                                      "Color": np.array([0.85, 0.325, 0.098]), "LineStyle": "--", "Marker": "square"})
    tx = node("text", {"String": "Speed (rpm)"})
    ty = node("text", {"String": "Torque (Nm)"})
    ax = node("axes", {"XLim": np.array([2000, 9000]), "YLim": np.array([0, 300])}, [l1, l2, tx, ty])
    sio.savemat(path, {"hgS_070000": node("figure", {"Color": np.array([1, 1, 1])}, [ax])})


def main():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        fig = d / "classic.fig"
        classic_fig(str(fig).replace(".fig", ".mat"))
        (d / "classic.mat").rename(fig)

        lines = fe.fig_lines(fig)
        check("fig_lines finds both named lines", set(lines) == {"FW trajectory", "MTPV trajectory"}, list(lines))
        axes = fe.axes_of(fig)
        check("axis labels and limits read", axes[0]["xlabel"] == "Speed (rpm)" and axes[0]["ylim"] == (0.0, 300.0),
              axes[0])
        check("line colour / style / marker kept", axes[0]["lines"][1]["ls"] == "--"
              and axes[0]["lines"][1]["marker"] == "square")

        out = d / "csv"
        check("lines → csv", fe.main(["lines", str(fig), str(out)]) == 0 and len(list(out.glob("*.csv"))) == 2)

        png = d / "replot.png"
        check("replot rc 0", fe.main(["replot", str(fig), str(png), "--slot", "3.94", "2.75"]) == 0)
        with Image.open(png) as im:
            w, h = im.size
        check("replot pixel size = slot × 300 dpi", abs(w - 1182) <= 2 and abs(h - 825) <= 2, (w, h))

        hg2 = d / "hg2.fig"
        sio.savemat(str(hg2).replace(".fig", ".mat"), {"hgS_080000": np.zeros((2, 2))})
        (d / "hg2.mat").rename(hg2)
        check("HG2-style file reports a format error (rc 2)", fe.main(["summary", str(hg2)]) == 2)

        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(png), Inches(1), Inches(1), Inches(3.94), Inches(2.75))
        deck = d / "deck.pptx"
        prs.save(deck)
        check("slot lists the placed picture", fe.main(["slot", str(deck), "1"]) == 0)
    failed = ok.count(False)
    print("all passed" if not failed else f"{failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
