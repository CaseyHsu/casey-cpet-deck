"""fig_extract.py — read MATLAB .fig files without MATLAB and re-plot them at slide-slot size.

cpet-figure-prep v1.2「既有 MATLAB 圖重繪」(ECCE 2026 lesson G10; from ECCE build/replot/figio.py).

  summary <file.fig>                      print the object tree (type, data shapes, names)
  lines   <file.fig> <out_dir>            dump every line to <out_dir>/ax<i>_<name>.csv
  slot    <deck.pptx> <slide>             list the pictures on a slide with their size in inches
                                          (use as --slot when replacing a picture 1:1)
  replot  <file.fig> <out.png> --slot W H [--title-off]
                                          redraw the line plots at W x H inches, 300 dpi,
                                          keeping each line's colour / style / marker / name

Rules carried over from ECCE:
  * figsize = the slide slot size in inches, so the PNG is placed without scaling
    and fonts stay at their real point size (9-11 pt);
  * keep the paper's colours (the .fig's own Color property) so slides match the paper
    (conventions §2.5.5 exception: grey-scale test + line style / marker double coding);
  * Arial with Liberation Sans / DejaVu fallback, white background, 300 dpi.

Only the classic MAT v5 .fig layout (hgS_* root holding handle-graphics structs) is decoded.
Figures saved by newer MATLAB (HG2) store MCOS objects that scipy cannot read; the tool says
so and exits 2 — export each line's XData/YData from MATLAB instead and plot it in the project.
Surface / contour plots are listed by `summary` but not re-plotted.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

import numpy as np

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass

LINE_TYPES = ("line", "graph2d.lineseries")


class FigFormatError(ValueError):
    pass


def _node(n) -> dict:
    d = {"type": str(n.type), "props": {}, "children": []}
    p = n.properties
    if hasattr(p, "_fieldnames"):
        for k in p._fieldnames:
            d["props"][k] = getattr(p, k)
    ch = n.children
    if isinstance(ch, np.ndarray):
        ch = list(ch.flat)
    elif ch is None or (hasattr(ch, "size") and ch.size == 0):
        ch = []
    else:
        ch = [ch]
    for c in ch:
        if hasattr(c, "type"):
            d["children"].append(_node(c))
    return d


def load(path: str | Path) -> dict:
    import scipy.io as sio

    m = sio.loadmat(str(path), struct_as_record=False, squeeze_me=True)
    roots = [k for k in m if k.startswith("hgS")]
    if not roots:
        keys = [k for k in m if not k.startswith("__")]
        raise FigFormatError(f"{Path(path).name}: no hgS_* root (keys {keys}); "
                             "not a classic MATLAB figure — export the line data from MATLAB instead")
    root = m[roots[0]]
    if not hasattr(root, "type"):
        raise FigFormatError(f"{Path(path).name}: {roots[0]} holds {type(root).__name__} "
                             f"{getattr(root, 'dtype', '')}, not a handle-graphics struct; "
                             "HG2 figure (newer MATLAB) — export each line's XData/YData from MATLAB instead")
    return _node(root)


def walk(d: dict, depth: int = 0, out=None):
    out = [] if out is None else out
    out.append((depth, d))
    for c in d["children"]:
        walk(c, depth + 1, out)
    return out


def summary(path) -> None:
    for dep, d in walk(load(path)):
        p, info = d["props"], []
        for k in ("XData", "YData", "ZData", "CData"):
            if k in p and hasattr(p[k], "shape"):
                info.append(f"{k}{np.shape(p[k])}")
        for k in ("DisplayName", "String", "Tag", "Color", "LineStyle", "Marker", "LevelList"):
            if k in p:
                info.append(f"{k}={str(p[k])[:60]}")
        print("  " * dep + d["type"], " ".join(info))


def axes_of(path) -> list[dict]:
    """Per axes: labels, limits and its line objects (x, y, style)."""
    tree = load(path)
    out = []
    for _, n in walk(tree):
        if n["type"] != "axes":
            continue
        ax = {"xlabel": "", "ylabel": "", "title": "", "xlim": None, "ylim": None, "lines": []}
        p = n["props"]
        for lim in ("XLim", "YLim"):
            if lim in p and np.size(p[lim]) == 2:
                ax[lim.lower()] = tuple(float(v) for v in np.ravel(p[lim]))
        for _, c in walk(n):
            cp = c["props"]
            if c["type"] == "text" and "String" in cp:
                # label text objects are children of the axes; their role is not stored
                # in old .fig files except via the parent's handles, so keep order
                ax.setdefault("_texts", []).append(str(cp["String"]))
            if c["type"].startswith(LINE_TYPES) and "XData" in cp and "YData" in cp:
                x = np.asarray(cp["XData"], float).ravel()
                y = np.asarray(cp["YData"], float).ravel()
                if x.size < 2:
                    continue
                ax["lines"].append({
                    "name": str(cp.get("DisplayName", "")) or None,
                    "x": x, "y": y,
                    "color": tuple(np.ravel(cp["Color"]).tolist()) if "Color" in cp else None,
                    "ls": str(cp.get("LineStyle", "-")),
                    "marker": str(cp.get("Marker", "none")),
                    "lw": float(cp.get("LineWidth", 1.0)),
                })
        texts = ax.pop("_texts", [])
        # Observed order in old .fig files: XLabel, YLabel[, ZLabel, Title]; empty labels are omitted
        if len(texts) >= 2:
            ax["xlabel"], ax["ylabel"] = texts[0], texts[1]
        if len(texts) >= 4:
            ax["title"] = texts[-1]
        if ax["lines"]:
            out.append(ax)
    return out


def fig_lines(path) -> dict[str, tuple[np.ndarray, np.ndarray]]:
    """{DisplayName: (x, y)} of every named line (ECCE replot helper)."""
    return {l["name"]: (l["x"], l["y"]) for ax in axes_of(path) for l in ax["lines"] if l["name"]}


def _safe(name: str) -> str:
    return re.sub(r"[^\w.-]+", "_", name).strip("_") or "line"


def cmd_lines(a) -> int:
    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    n = 0
    for i, ax in enumerate(axes_of(a.fig), 1):
        for j, l in enumerate(ax["lines"], 1):
            fn = out / f"ax{i}_{_safe(l['name'] or f'line{j}')}.csv"
            with open(fn, "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["x", "y"])
                w.writerows(zip(l["x"], l["y"]))
            n += 1
    print(f"wrote {n} csv files to {out}")
    return 0 if n else 1


_MARK = {"none": None, "x": "x", "o": "o", "+": "+", "*": "*", "square": "s", "diamond": "D",
         "^": "^", "v": "v", ".": "."}
_LS = {"-": "-", "--": "--", ":": ":", "-.": "-.", "none": "none"}


def setup_style():
    import matplotlib as mpl

    mpl.rcParams.update({
        "font.family": ["Arial", "Liberation Sans", "DejaVu Sans"], "font.size": 10,
        "axes.labelsize": 11, "xtick.labelsize": 10, "ytick.labelsize": 10, "legend.fontsize": 9,
        "axes.linewidth": 1.0, "grid.color": "#D9D9D9", "grid.linewidth": 0.6, "legend.frameon": False,
        "savefig.dpi": 300,
    })


def cmd_replot(a) -> int:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    setup_style()
    axes = axes_of(a.fig)
    if not axes:
        print("no line plots found")
        return 1
    w, h = a.slot
    fig, axs = plt.subplots(len(axes), 1, figsize=(w, h), squeeze=False)
    for ax_d, ax in zip(axes, axs[:, 0]):
        for l in ax_d["lines"]:
            ax.plot(l["x"], l["y"], color=l["color"], linestyle=_LS.get(l["ls"], "-"),
                    marker=_MARK.get(l["marker"]), markersize=4, markevery=max(1, len(l["x"]) // 20),
                    linewidth=max(1.2, l["lw"]), label=l["name"])
        ax.set_xlabel(ax_d["xlabel"]); ax.set_ylabel(ax_d["ylabel"])
        if ax_d["title"] and not a.title_off:
            ax.set_title(ax_d["title"])
        if ax_d["xlim"]:
            ax.set_xlim(ax_d["xlim"])
        if ax_d["ylim"]:
            ax.set_ylim(ax_d["ylim"])
        ax.grid(True)
        if any(l["name"] for l in ax_d["lines"]):
            ax.legend(loc="best")
    fig.tight_layout()
    fig.savefig(a.out, dpi=300, facecolor="white")
    print(f"saved {a.out} ({w} x {h} in @300 dpi)")
    return 0


def cmd_slot(a) -> int:
    """Pictures on one slide with their placed size: replacing one 1:1 means figsize = this size."""
    from pptx import Presentation

    prs = Presentation(a.deck)
    if not 1 <= a.slide <= len(prs.slides):
        print(f"slide {a.slide} out of range (1–{len(prs.slides)})", file=sys.stderr)
        return 2
    pics = [sh for sh in prs.slides[a.slide - 1].shapes if sh.shape_type == 13]
    pics.sort(key=lambda sh: sh.width * sh.height, reverse=True)   # 內容圖在前，模板 logo 沉到後面
    n = len(pics)
    for sh in pics:
        print(f"{sh.name}: --slot {sh.width / 914400:.2f} {sh.height / 914400:.2f}"
              f"  (at {sh.left / 914400:.2f}, {sh.top / 914400:.2f} in)")
    if not n:
        print("no pictures on this slide")
    return 0 if n else 1


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("summary"); s.add_argument("fig")
    s.set_defaults(fn=lambda a: summary(a.fig) or 0)
    s = sub.add_parser("lines"); s.add_argument("fig"); s.add_argument("out_dir"); s.set_defaults(fn=cmd_lines)
    s = sub.add_parser("slot"); s.add_argument("deck"); s.add_argument("slide", type=int); s.set_defaults(fn=cmd_slot)
    s = sub.add_parser("replot"); s.add_argument("fig"); s.add_argument("out")
    s.add_argument("--slot", nargs=2, type=float, required=True, metavar=("W_IN", "H_IN"))
    s.add_argument("--title-off", action="store_true")
    s.set_defaults(fn=cmd_replot)
    a = ap.parse_args(argv)
    try:
        return a.fn(a)
    except FigFormatError as e:
        print("format:", e, file=sys.stderr)
        return 2
    except (OSError, ValueError, NotImplementedError) as e:
        print("error:", e, file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
